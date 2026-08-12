"""Lettura del piano particellare e costruzione incrementale dell'output storico."""

from __future__ import annotations

import json
import re
import shutil
import unicodedata
from copy import copy
from dataclasses import asdict, dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from config.settings import Settings


def _text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _header(value) -> str:
    normalized = unicodedata.normalize("NFKD", _text(value).lower())
    ascii_text = "".join(char for char in normalized if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", ascii_text)


def _decimal(value) -> Decimal:
    raw = _text(value).replace(".", "").replace(",", ".")
    if not raw:
        return Decimal(0)
    try:
        return Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError(f"Valore di superficie non numerico: {value!r}") from exc


def calculate_square_metres(ha, are, ca):
    """Converte ettari, are e centiare in metri quadrati."""
    result = (_decimal(ha) * 10_000) + (_decimal(are) * 100) + _decimal(ca)
    return int(result) if result == result.to_integral_value() else float(result)


@dataclass(frozen=True)
class HistoricalJob:
    order: str
    source_row: int
    comune: str
    foglio: str
    particella: str
    sezione: str = ""

    @property
    def cache_key(self) -> str:
        pieces = (self.order, self.comune, self.foglio, self.particella)
        clean = [re.sub(r"[^A-Za-z0-9.-]+", "_", part).strip("_") for part in pieces]
        return "STORICA_" + "_".join(clean).upper()

    @property
    def pdf_name(self) -> str:
        pieces = (self.order, self.comune, self.foglio, self.particella)
        clean = [re.sub(r'[<>:"/\\|?*]+', "_", part).strip() for part in pieces]
        return "_".join(clean) + ".pdf"


@dataclass
class PropertyResult:
    sub: str = ""
    ha: str = ""
    are: str = ""
    ca: str = ""
    qualita: str = ""
    reddito_dominicale: str = ""
    reddito_agrario: str = ""
    intestatari: list[str] = field(default_factory=list)

    @property
    def mq(self):
        if not any(_text(value) for value in (self.ha, self.are, self.ca)):
            return ""
        return calculate_square_metres(self.ha, self.are, self.ca)


class HistoricalWorkbook:
    """Mantiene il sorgente intatto e rigenera l'output dai risultati in cache."""

    REQUIRED_HEADERS = {
        "ordine": {"nordine", "ordine"},
        "proprietario": {"dittecatastalioproprietarioedindirizzo"},
        "comune": {"comune"},
        "foglio": {"foglio"},
        "particella": {"particelle", "particella"},
        "sub": {"sub"},
        "ha": {"ha"},
        "are": {"are"},
        "ca": {"cent", "ca", "centiare"},
        "mq": {"mq"},
        "qualita": {"qualitaterrenocategoria", "qualita"},
        "reddito_dominicale": {"redditodom", "redditodominicale"},
        "reddito_agrario": {"redditoagr", "redditoagrario"},
    }

    def __init__(self, source: str | Path, output: str | Path | None = None):
        self.source = Path(source)
        self.output = Path(output or Settings.DESTINATION_EXCEL)
        self.output.parent.mkdir(parents=True, exist_ok=True)
        Settings.HISTORICAL_DIR.mkdir(parents=True, exist_ok=True)
        Settings.HISTORICAL_DATA_CACHE_DIR.mkdir(parents=True, exist_ok=True)
        self.sheet_name, self.header_row, self.columns = self._inspect_source()

    def _inspect_source(self):
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        try:
            sheet = workbook.active
            for row_number, row in enumerate(sheet.iter_rows(), start=1):
                normalized = {_header(cell.value): cell.column for cell in row if _text(cell.value)}
                if "comune" in normalized and "foglio" in normalized and (
                    "particelle" in normalized or "particella" in normalized
                ):
                    columns = {}
                    for logical, aliases in self.REQUIRED_HEADERS.items():
                        for alias in aliases:
                            if alias in normalized:
                                columns[logical] = normalized[alias]
                                break
                    missing = set(self.REQUIRED_HEADERS) - set(columns)
                    if missing:
                        raise ValueError(f"Colonne mancanti nel file Excel: {', '.join(sorted(missing))}")
                    return sheet.title, row_number, columns
        finally:
            workbook.close()
        raise ValueError("Impossibile trovare l'intestazione del piano particellare.")

    @staticmethod
    def section_for(comune: str) -> str:
        return "" if "SERMIDE" in comune.upper() else "B"

    def jobs(self) -> list[HistoricalJob]:
        workbook = load_workbook(self.source, read_only=True, data_only=True)
        try:
            sheet = workbook[self.sheet_name]
            jobs = []
            for row_number in range(self.header_row + 1, sheet.max_row + 1):
                order = _text(sheet.cell(row_number, self.columns["ordine"]).value)
                comune = _text(sheet.cell(row_number, self.columns["comune"]).value)
                foglio = _text(sheet.cell(row_number, self.columns["foglio"]).value)
                particella = _text(sheet.cell(row_number, self.columns["particella"]).value)
                if not (order and comune and foglio and particella):
                    continue
                jobs.append(HistoricalJob(
                    order=order,
                    source_row=row_number,
                    comune=comune,
                    foglio=foglio,
                    particella=particella,
                    sezione=self.section_for(comune),
                ))
            return jobs
        finally:
            workbook.close()

    def _result_path(self, job: HistoricalJob) -> Path:
        return Settings.HISTORICAL_DATA_CACHE_DIR / f"{job.cache_key}.json"

    def has_result(self, job: HistoricalJob) -> bool:
        return self._result_path(job).exists()

    def save_result(self, job: HistoricalJob, result: PropertyResult) -> None:
        path = self._result_path(job)
        temporary = path.with_suffix(".tmp")
        temporary.write_text(json.dumps(asdict(result), ensure_ascii=False, indent=2), encoding="utf-8")
        temporary.replace(path)

    def load_result(self, job: HistoricalJob) -> PropertyResult | None:
        path = self._result_path(job)
        if not path.exists():
            return None
        return PropertyResult(**json.loads(path.read_text(encoding="utf-8")))

    @staticmethod
    def _copy_row_style(sheet, source_row: int, destination_row: int) -> None:
        sheet.row_dimensions[destination_row].height = sheet.row_dimensions[source_row].height
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row, column)
            destination = sheet.cell(destination_row, column)
            if source.has_style:
                destination._style = copy(source._style)
            destination.number_format = source.number_format
            destination.alignment = copy(source.alignment)
            destination.protection = copy(source.protection)

    def build_destination(self, jobs: list[HistoricalJob]) -> Path:
        shutil.copy2(self.source, self.output)
        workbook = load_workbook(self.output)
        sheet = workbook[self.sheet_name]

        # Inserendo dal basso, source_row continua a riferirsi alla riga originale.
        for job in reversed(jobs):
            result = self.load_result(job)
            if result is None:
                continue

            row = job.source_row
            values = {
                "proprietario": result.intestatari[0] if result.intestatari else "",
                "sub": result.sub,
                "ha": result.ha,
                "are": result.are,
                "ca": result.ca,
                "mq": result.mq,
                "qualita": result.qualita,
                "reddito_dominicale": result.reddito_dominicale,
                "reddito_agrario": result.reddito_agrario,
            }
            for logical, value in values.items():
                sheet.cell(row, self.columns[logical]).value = value

            extra_owners = result.intestatari[1:]
            if extra_owners:
                sheet.insert_rows(row + 1, amount=len(extra_owners))
                for offset, owner in enumerate(extra_owners, start=1):
                    destination_row = row + offset
                    self._copy_row_style(sheet, row, destination_row)
                    for column in range(1, sheet.max_column + 1):
                        sheet.cell(destination_row, column).value = None
                    sheet.cell(destination_row, self.columns["ordine"]).value = self._letter(offset)
                    sheet.cell(destination_row, self.columns["proprietario"]).value = owner

        workbook.save(self.output)
        return self.output

    @staticmethod
    def _letter(position: int) -> str:
        result = ""
        while position:
            position, remainder = divmod(position - 1, 26)
            result = chr(97 + remainder) + result
        return result

    def publish_pdf(self, job: HistoricalJob, cached_pdf: Path) -> Path:
        destination = Settings.HISTORICAL_DIR / job.pdf_name
        shutil.copy2(cached_pdf, destination)
        return destination
