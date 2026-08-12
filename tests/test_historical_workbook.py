import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from config.settings import Settings
from core.historical_workbook import (
    HistoricalWorkbook,
    PropertyResult,
    calculate_square_metres,
)


HEADERS = (
    "n.\nOrdine",
    "Ditte catastali o Proprietario ed indirizzo",
    "Comune",
    "foglio",
    "particelle",
    "sub",
    "ha",
    "are",
    "cent.",
    "MQ",
    "Qualità terreno\ncategoria",
    "Reddito Dom.",
    "Reddito Agr.",
)


class HistoricalWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)
        self.source = self.root / "input.xlsx"
        self.output = self.root / "output" / "destinazione.xlsx"
        self.cache_dir = self.root / "cache" / "dati_storici"
        self.historical_dir = self.root / "output" / "storica"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Foglio1"
        for column, value in enumerate(HEADERS, start=4):
            sheet.cell(7, column).value = value
        sheet.append([])
        for row, values in enumerate((
            (1, None, "Borgocarbonara", 7, 127),
            (2, None, "Borgocarbonara", 7, 119),
            (3, None, "Sermide e Felonica", 1, 110),
        ), start=8):
            for column, value in enumerate(values, start=4):
                sheet.cell(row, column).value = value
        workbook.save(self.source)

        self.settings_patch = (
            patch.object(Settings, "HISTORICAL_DATA_CACHE_DIR", self.cache_dir),
            patch.object(Settings, "HISTORICAL_DIR", self.historical_dir),
        )
        for item in self.settings_patch:
            item.start()

    def tearDown(self):
        for item in reversed(self.settings_patch):
            item.stop()
        self.temp_dir.cleanup()

    def test_jobs_apply_section_b_except_sermide(self):
        manager = HistoricalWorkbook(self.source, self.output)
        jobs = manager.jobs()

        self.assertEqual([job.order for job in jobs], ["1", "2", "3"])
        self.assertEqual([job.sezione for job in jobs], ["B", "B", ""])
        self.assertEqual(jobs[2].pdf_name, "3_Sermide e Felonica_1_110.pdf")

    def test_build_preserves_source_and_inserts_alphabetical_owner_rows(self):
        manager = HistoricalWorkbook(self.source, self.output)
        jobs = manager.jobs()
        manager.save_result(jobs[0], PropertyResult(
            sub="",
            ha="0",
            are="1",
            ca="37",
            qualita="ENTE URBANO",
            intestatari=["Intestatario uno", "Intestatario due", "Intestatario tre"],
        ))
        manager.save_result(jobs[1], PropertyResult(
            ha="1", are="2", ca="3", intestatari=["Altro proprietario"]
        ))

        manager.build_destination(jobs)

        source_sheet = load_workbook(self.source, data_only=True).active
        self.assertIsNone(source_sheet["E8"].value)
        destination = load_workbook(self.output, data_only=True).active
        self.assertEqual(destination["D8"].value, 1)
        self.assertEqual(destination["E8"].value, "Intestatario uno")
        self.assertEqual(destination["M8"].value, 137)
        self.assertEqual(destination["N8"].value, "ENTE URBANO")
        self.assertEqual((destination["D9"].value, destination["E9"].value), ("a", "Intestatario due"))
        self.assertEqual((destination["D10"].value, destination["E10"].value), ("b", "Intestatario tre"))
        self.assertEqual((destination["D11"].value, destination["E11"].value), (2, "Altro proprietario"))
        self.assertEqual(destination["M11"].value, 10203)

    def test_square_metre_conversion_accepts_italian_decimals(self):
        self.assertEqual(calculate_square_metres("0", "1", "37"), 137)
        self.assertEqual(calculate_square_metres("1", "2", "3"), 10203)
        self.assertEqual(calculate_square_metres("0", "0", "1,5"), 1.5)


if __name__ == "__main__":
    unittest.main()
