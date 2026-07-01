# (Fase 6) Generazione del report unificato delle eccezioni

# utils/reporters.py — Report unificato delle eccezioni catastali e di parsing
import pandas as pd
from datetime import datetime
from config.settings import Settings
from utils.logger import get_logger
import openpyxl
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

logger = get_logger("Reporters")

def generate_report(invalid_rows: list, scrape_errors: list, original_headers: list = None) -> None:
    """
    Unisce le righe scartate dal parser Excel e gli errori rilevati su Sister
    in un unico file Excel riassuntivo inserito nella cartella di output,
    con disposizione ordinata delle colonne ed evidenziazione bordeaux.
    """
    logger.info("Avvio generazione del report unificato delle eccezioni...")
    all_errors = []

    for err in invalid_rows + scrape_errors:
        if isinstance(err, str):
            all_errors.append({
                "TIPO ERRORE": err,
                "RIGA EXCEL": "N/D"
            })
        elif isinstance(err, dict):
            new_err = {}
            for k, v in err.items():
                if k in ("TIPO_ERRORE", "TIPO ERRORE"):
                    new_err["TIPO ERRORE"] = v
                elif k in ("RIGA_EXCEL", "RIGA EXCEL"):
                    new_err["RIGA EXCEL"] = v
                else:
                    new_err[k] = v
            # Assicuriamo la presenza delle chiavi minime
            if "TIPO ERRORE" not in new_err:
                new_err["TIPO ERRORE"] = new_err.get("errore", "Dato non valido o errore sconosciuto")
            if "RIGA EXCEL" not in new_err:
                new_err["RIGA EXCEL"] = new_err.get("riga", "N/D")
            all_errors.append(new_err)
        else:
            riga_num = getattr(err, "riga_excel", "N/D")
            motivo = getattr(err, "motivo", "Dato non valido")
            dati_orig = getattr(err, "dati_originali", {})
            
            err_dict = {
                "TIPO ERRORE": motivo,
                "RIGA EXCEL": riga_num,
            }
            if isinstance(dati_orig, dict):
                err_dict.update(dati_orig)
            all_errors.append(err_dict)

    # Se non ci sono anomalie, saltiamo
    if not all_errors:
        logger.info("✅ Complimenti! Nessun errore o scarto registrato. Salto la generazione del report.")
        return

    # Ordinamento cronologico per "RIGA EXCEL"
    def get_sort_key(item):
        val = item.get("RIGA EXCEL", "N/D")
        try:
            return (0, int(val))
        except (ValueError, TypeError):
            return (1, str(val))

    all_errors.sort(key=get_sort_key)

    # Definizione dell'ordine finale delle colonne
    # DOCUMENTO_IN_CACHE è valorizzata dall'orchestratore: 'Sì' / 'No' / assente per errori di parsing
    ordered_cols = ["TIPO ERRORE", "RIGA EXCEL", "DOCUMENTO_IN_CACHE"]
    if original_headers:
        for col in original_headers:
            if col not in ordered_cols:
                ordered_cols.append(col)

    # Creazione DataFrame e riordino colonne
    df = pd.DataFrame(all_errors)
    # Aggiungi eventuali colonne extra non comprese in original_headers
    for col in df.columns:
        if col not in ordered_cols:
            ordered_cols.append(col)
            
    df = df.reindex(columns=ordered_cols)

    # Percorso del file di output
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_file = Settings.OUTPUT_DIR / f"report_eccezioni_{ts}.xlsx"

    try:
        out_file.parent.mkdir(parents=True, exist_ok=True)
        
        # Scrittura del file Excel
        df.to_excel(str(out_file), index=False)

        # Formattazione visiva tramite openpyxl
        wb = openpyxl.load_workbook(str(out_file))
        ws = wb.active

        # Stile Bordeaux (rosso scuro bordeaux: "800020") e testo bianco grassetto per Colonna A (TIPO ERRORE)
        fill_bordeaux = PatternFill(start_color="800020", end_color="800020", fill_type="solid")
        font_white = Font(name="Calibri", size=11, bold=True, color="FFFFFF")

        for r_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=r_idx, column=1)
            cell.fill = fill_bordeaux
            cell.font = font_white

        # Stile Header professionale (sfondo grigio chiaro e testo in grassetto)
        fill_header = PatternFill(start_color="EAEAEA", end_color="EAEAEA", fill_type="solid")
        font_header = Font(name="Calibri", size=11, bold=True, color="000000")
        for c_idx in range(1, ws.max_column + 1):
            cell = ws.cell(row=1, column=c_idx)
            cell.fill = fill_header
            cell.font = font_header

        # Auto-fit larghezza colonne
        for col in ws.columns:
            max_len = 0
            col_letter = get_column_letter(col[0].column)
            for cell in col:
                val_str = str(cell.value or '')
                if len(val_str) > max_len:
                    max_len = len(val_str)
            # Regola larghezza min 10 e max 50
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 10), 50)

        wb.save(str(out_file))
        logger.info(f"✅ Report delle eccezioni salvato con successo in: {out_file.name}")
        logger.info(f"   📊 Riepilogo anomalie -> Totale: {len(all_errors)}")

    except Exception as e:
        logger.error(f"Errore critico durante la scrittura del file Excel di report: {e}")
